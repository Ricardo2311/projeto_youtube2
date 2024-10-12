from selenium import webdriver
from selenium.webdriver.common.by import By
import openpyxl
from docx import Document
from time import sleep
import os
# Pedir informacao do usuario
email = input('Digite o seu email: ')
senha = input('Digite a sua senha: ')
driver = webdriver.Chrome()
driver.get('https://contabil-devaprender.netlify.app')
def resto():
    email_registro = driver.find_element(By.XPATH, "//input[@type='email']")
    email_registro.send_keys(email)
    sleep(1)
    senha_registro = driver.find_element(By.XPATH, "//input[@type='password']")
    senha_registro.send_keys(senha)
    sleep(1)
    botao_login = driver.find_element(By.XPATH, "//button[@type='submit']")
    botao_login.click()
    sleep(3)

    # Clicar em Cadastrar Balanco Patrimonia
    botao_cadastrar_balanco_patrimonial = driver.find_elements(By.XPATH, "//a[@class='btn btn-primary mt-auto']")[0]
    botao_cadastrar_balanco_patrimonial.click()
resto()
# Verificar o tipo de arquivo
#WORD
def ler_arquivos_word(caminho_arquivo_word):

    doc = Document(caminho_arquivo_word)
    tabelas = doc.tables
    ativo_circulante = ''
    caixa_equivalentes = ''
    contas_receber = ''
    estoques = ''
    ativo_nao_circulante = ''
    imobilizado = ''
    intangivel = ''
    total_ativo = ''
    for tabela in tabelas:
        for linha in tabela.rows:
            if 'Ativo Circulante' in linha.cells[0].text.strip():
                ativo_circulante = linha.cells[1].text.strip()

            elif 'Caixa e Equivalentes' in linha.cells[0].text.strip():
                caixa_equivalentes = linha.cells[1].text.strip()

            elif 'Contas a Receber' in linha.cells[0].text.strip():
                contas_receber = linha.cells[1].text.strip()

            elif 'Estoques' in linha.cells[0].text.strip():
                estoques = linha.cells[1].text.strip()
            elif 'Ativo Não Circulante' in linha.cells[0].text.strip():
                ativo_nao_circulante = linha.cells[1].text.strip()

            elif 'Imobilizado' in linha.cells[0].text.strip():
                imobilizado = linha.cells[1].text.strip()

            elif 'Intangível' in linha.cells[0].text.strip():
                intangivel = linha.cells[1].text.strip()

            elif 'Total do Ativo' in linha.cells[0].text.strip():
                total_ativo = linha.cells[1].text.strip()
    sleep(1)
    campo_ativo_circulante = driver.find_element(By.XPATH, "//input[@id='ativo_circulante']")
    campo_caixa_equivalentes = driver.find_element(By.XPATH, "//input[@id='caixa_equivalentes']")
    campo_contas_receber = driver.find_element(By.XPATH, "//input[@id='contas_receber']")
    campo_estoques = driver.find_element(By.XPATH, "//input[@id='estoques']")
    campo_ativo_nao_circulante = driver.find_element(By.XPATH, "//input[@id='ativo_nao_circulante']")
    campo_imobilizado = driver.find_element(By.XPATH, "//input[@id='imobilizado']")
    campo_intangivel = driver.find_element(By.XPATH, "//input[@id='intangivel']")
    campo_total_ativo = driver.find_element(By.XPATH, "//input[@id='total_ativo']")
    sleep(1)
    campo_ativo_circulante.send_keys(ativo_circulante)
    sleep(1)
    campo_caixa_equivalentes.send_keys(caixa_equivalentes)
    sleep(1)
    campo_contas_receber.send_keys(contas_receber)
    sleep(1)
    campo_estoques.send_keys(estoques)
    sleep(1)
    campo_ativo_nao_circulante.send_keys(ativo_nao_circulante)
    sleep(1)
    campo_imobilizado.send_keys(imobilizado)
    sleep(1)
    campo_intangivel.send_keys(intangivel)
    sleep(1)
    campo_total_ativo.send_keys(total_ativo)
    sleep(1)
    botao_cadastrar = driver.find_element(By.XPATH, "//button[@type='submit']")
    botao_cadastrar.click()
caminho_atual = r'C:\Users\rgeba\OneDrive\Documentos\CURSO DEV APRENDER\Projeto Destrava Web\Projeto Youtube #3'
for arquivo in os.listdir(caminho_atual):
    if arquivo.endswith('.docx'):
        caminho_arquivo_word = os.path.join(caminho_atual,arquivo)
        ler_arquivos_word(caminho_arquivo_word)
#EXCEL
def ler_arquivos_excel(caminho_arquivo_excel):
    planilha = openpyxl.load_workbook(caminho_arquivo_excel)
    try:
        pagina = planilha['Sheet1']
    except:
        pagina = planilha['Planilha1']
    for linha in pagina.iter_rows(min_row=2,values_only=True):
        ativo_circulante = linha[0]
        caixa_equivalentes = linha[1]
        contas_receber = linha[2]
        estoques = linha[3]
        ativo_nao_circulante = linha[4]
        imobilizado = linha[5]
        intangivel = linha[6]
        total_ativo = linha[7]
        sleep(1)
        campo_ativo_circulante = driver.find_element(By.XPATH, "//input[@id='ativo_circulante']")
        campo_caixa_equivalentes = driver.find_element(By.XPATH, "//input[@id='caixa_equivalentes']")
        campo_contas_receber = driver.find_element(By.XPATH, "//input[@id='contas_receber']")
        campo_estoques = driver.find_element(By.XPATH, "//input[@id='estoques']")
        campo_ativo_nao_circulante = driver.find_element(By.XPATH, "//input[@id='ativo_nao_circulante']")
        campo_imobilizado = driver.find_element(By.XPATH, "//input[@id='imobilizado']")
        campo_intangivel = driver.find_element(By.XPATH, "//input[@id='intangivel']")
        campo_total_ativo = driver.find_element(By.XPATH, "//input[@id='total_ativo']")
        sleep(1)
        campo_ativo_circulante.send_keys(ativo_circulante)
        sleep(1)
        campo_caixa_equivalentes.send_keys(caixa_equivalentes)
        sleep(1)
        campo_contas_receber.send_keys(contas_receber)
        sleep(1)
        campo_estoques.send_keys(estoques)
        sleep(1)
        campo_ativo_nao_circulante.send_keys(ativo_nao_circulante)
        sleep(1)
        campo_imobilizado.send_keys(imobilizado)
        sleep(1)
        campo_intangivel.send_keys(intangivel)
        sleep(1)
        campo_total_ativo.send_keys(total_ativo)
        sleep(1)
        botao_cadastrar = driver.find_element(By.XPATH, "//button[@type='submit']")
        botao_cadastrar.click()
for arquivo in os.listdir(caminho_atual):
    if arquivo.endswith('.xlsx'):
        caminho_arquivo_excel = os.path.join(caminho_atual,arquivo)
        ler_arquivos_excel(caminho_arquivo_excel)
